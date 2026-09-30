"""Bulk customer import from an .xlsx template -- see
docs/superpowers/specs/2026-09-30-customer-import-wizard-design.md.

Migration mode: customers arrive with their current paid-through expiry and
nothing is billed, linked to Whish, or sent on WhatsApp. Registered from
app.py via register_customer_import_routes(app, appmod) so app.py doesn't
grow further."""
import io
from functools import wraps

from flask import jsonify, request, send_file
from flask_jwt_extended import verify_jwt_in_request
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from tenancy import tenant_query

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
SHEET_NAME = 'Customers'
XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# (header, required, instruction). Order = template column order.
COLUMNS = [
    ('name', True, 'Customer full name (max 100 characters).'),
    ('phone', True, 'Phone number, kept as text so leading zeros survive (max 20).'),
    ('address', True, 'Address (max 200 characters).'),
    ('plan', True, 'Subscription plan name. Pick from the list, or type a new name -- '
                   'the wizard will ask for its price before importing.'),
    ('expiry_date', True, 'Date the customer is paid through (YYYY-MM-DD). '
                          'Billing resumes from this date.'),
    ('active', False, 'yes / no. Default yes. Inactive customers are never billed.'),
    ('sector', False, 'Area / sector. New names are created automatically.'),
    ('reseller', False, 'Reseller name, exactly as in ServiceBills. Blank for direct customers.'),
    ('discount', False, 'Discount per billing cycle (number, default 0).'),
    ('cost_override', False, 'Your cost for this customer, if different from the plan cost.'),
    ('opening_balance', False, 'Amount the customer owes you today (default 0). '
                               'Recorded as one unpaid payment.'),
    ('upstream_provider', False, 'Upstream provider name, exactly as in ServiceBills.'),
    ('upstream_username', False, 'Username on the upstream provider portal.'),
    ('network_device', False, 'Network device (MikroTik) name, exactly as in ServiceBills.'),
    ('pppoe_username', False, 'PPPoE username on that network device.'),
    ('onu_mac', False, 'ONU MAC address (aa:bb:cc:dd:ee:ff).'),
    ('cpe_mac', False, "Customer's own router MAC. One router per customer."),
    ('notes', False, 'Free text notes.'),
    ('whatsapp_enabled', False, 'yes / no. Default yes.'),
]
HEADERS = [c[0] for c in COLUMNS]
REQUIRED = [c[0] for c in COLUMNS if c[1]]
TEXT_COLUMNS = ('phone', 'upstream_username', 'pppoe_username', 'onu_mac', 'cpe_mac')
# Columns with a dropdown fed from the hidden Lists sheet. True = typing a
# value that isn't in the list is allowed (plan/sector get created on import).
LOOKUP_COLUMNS = {'plan': True, 'sector': True, 'reseller': False,
                  'upstream_provider': False, 'network_device': False}
MIGRATION_NOTES = [
    'How the import works:',
    '- Nothing is billed on import. expiry_date is the date the customer is paid through; '
    'the normal billing run charges the next cycle on that date.',
    '- A customer whose expiry_date has already passed (and is active) will be charged on the '
    'next billing run, exactly like any overdue customer.',
    '- opening_balance is what the customer owes you today; it becomes one unpaid payment.',
    '- No WhatsApp messages are sent to imported customers.',
    '- Rows that repeat a customer (same name and phone) already in ServiceBills are skipped, '
    'so you can safely upload a corrected file again.',
]


class ImportFileError(ValueError):
    """The upload as a whole is unusable (wrong type, headers, size)."""


def _lookup_names(appmod):
    """Tenant-scoped, sorted, de-duplicated names for each dropdown column."""
    models = {'plan': appmod.SubscriptionPlan, 'sector': appmod.Sector,
              'reseller': appmod.Reseller, 'upstream_provider': appmod.UpstreamProvider,
              'network_device': appmod.NetworkDevice}
    return {key: sorted({r.name for r in tenant_query(model).all() if r.name})
            for key, model in models.items()}


def build_template(appmod):
    """Return the current tenant's import template as .xlsx bytes."""
    lists = _lookup_names(appmod)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    letters = {}
    for idx, header in enumerate(HEADERS, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        letters[header] = letter
        ws.column_dimensions[letter].width = max(14, len(header) + 4)
        if header in TEXT_COLUMNS:
            ws.column_dimensions[letter].number_format = '@'
    last_row = MAX_ROWS + 1

    ins = wb.create_sheet('Instructions')
    ins.append(['Column', 'Required', 'What to enter'])
    for cell in ins[1]:
        cell.font = Font(bold=True)
    for header, required, help_text in COLUMNS:
        ins.append([header, 'yes' if required else '', help_text])
    ins.append([])
    for line in MIGRATION_NOTES:
        ins.append([line])
    ins.column_dimensions['A'].width = 20
    ins.column_dimensions['C'].width = 90

    lists_ws = wb.create_sheet('Lists')
    for col, (key, allow_new) in enumerate(LOOKUP_COLUMNS.items(), start=1):
        values = lists[key]
        lists_ws.cell(row=1, column=col, value=key)
        for i, value in enumerate(values, start=2):
            lists_ws.cell(row=i, column=col, value=value)
        if not values:
            continue
        src = lists_ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type='list',
                            formula1=f'=Lists!${src}$2:${src}${len(values) + 1}',
                            allow_blank=True, showErrorMessage=not allow_new)
        ws.add_data_validation(dv)
        dv.add(f'{letters[key]}2:{letters[key]}{last_row}')
    for key in ('active', 'whatsapp_enabled'):
        dv = DataValidation(type='list', formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{letters[key]}2:{letters[key]}{last_row}')
    lists_ws.sheet_state = 'hidden'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def read_workbook(file_storage):
    """Parse an uploaded template into a list of {header: value} dicts, each
    with '_row' = its 1-based sheet row. Fully blank rows are skipped.
    Raises ImportFileError when the file as a whole is unusable."""
    if file_storage is None or not getattr(file_storage, 'filename', ''):
        raise ImportFileError('No file uploaded.')
    data = file_storage.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError('File is larger than 5 MB.')
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError('Could not read the file. Upload the .xlsx template (Excel format).')
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ImportFileError(f"The workbook has no '{SHEET_NAME}' sheet. "
                                  f"Start from the downloaded template.")
        rows = wb[SHEET_NAME].iter_rows(values_only=True)
        header_row = next(rows, None) or ()
        headers = [str(h).strip().lower() if h is not None else '' for h in header_row]
        missing = [h for h in REQUIRED if h not in headers]
        if missing:
            raise ImportFileError('Missing required column(s): ' + ', '.join(missing))
        index = {h: i for i, h in enumerate(headers) if h in HEADERS}
        out = []
        for row_number, values in enumerate(rows, start=2):
            if not values or all(_is_blank(v) for v in values):
                continue
            if len(out) >= MAX_ROWS:
                raise ImportFileError(f'Too many rows: the limit is {MAX_ROWS} customers per import.')
            record = {h: (values[i] if i < len(values) else None) for h, i in index.items()}
            record['_row'] = row_number
            out.append(record)
    finally:
        wb.close()
    if not out:
        raise ImportFileError('The Customers sheet has no data rows.')
    return out


def register_customer_import_routes(app, appmod):
    def import_admin(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            verify_jwt_in_request()
            if 'admin' not in appmod._jwt_roles():
                return jsonify(msg="Admins only!"), 403
            return fn(*args, **kwargs)
        return wrapper

    @app.route('/api/customers/import/template', methods=['GET'])
    @import_admin
    def customer_import_template():
        return send_file(io.BytesIO(build_template(appmod)), mimetype=XLSX_MIME,
                         as_attachment=True, download_name='customer-import-template.xlsx')
