"""Customer import wizard -- see
docs/superpowers/specs/2026-09-30-customer-import-wizard-design.md."""
import io
import json
from datetime import datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import make_tenant, auth_headers
import customer_import as ci

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _plan(client, headers, name='Basic', price=20, cycle='monthly'):
    r = client.post('/api/subscription_plans', headers=headers,
                    json={'name': name, 'price': price, 'billing_cycle': cycle})
    assert r.status_code in (200, 201), r.get_json()
    return name


def _xlsx(rows, headers=None):
    """Build an upload workbook: `rows` is a list of dicts keyed by header."""
    headers = headers or ci.HEADERS
    wb = Workbook()
    ws = wb.active
    ws.title = 'Customers'
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _future(days=20):
    return (datetime.utcnow() + timedelta(days=days)).strftime('%Y-%m-%d')


def _row(**kw):
    base = {'name': 'Ali', 'phone': '03123456', 'address': 'Beirut',
            'plan': 'Basic', 'expiry_date': _future()}
    base.update(kw)
    return base


def test_template_requires_admin(client):
    make_tenant(client, 'Biz', 'owner')
    cashier = auth_headers(client, 'cash1', role='cashier')
    r = client.get('/api/customers/import/template', headers=cashier)
    assert r.status_code == 403


def test_template_lists_tenant_plans_only(client):
    a = make_tenant(client, 'A', 'a_admin')
    _plan(client, a, 'Fiber A')
    b = make_tenant(client, 'B', 'b_admin')
    _plan(client, b, 'Fiber B')

    r = client.get('/api/customers/import/template', headers=b)
    assert r.status_code == 200
    assert r.mimetype == XLSX
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ['Customers', 'Instructions', 'Lists']
    assert [c.value for c in wb['Customers'][1]] == ci.HEADERS
    lists_values = {c.value for col in wb['Lists'].iter_cols() for c in col}
    assert 'Fiber B' in lists_values
    assert 'Fiber A' not in lists_values
    assert wb['Lists'].sheet_state == 'hidden'


class _Upload:
    """Minimal stand-in for werkzeug FileStorage."""
    def __init__(self, buf, filename='c.xlsx'):
        self._buf = buf
        self.filename = filename

    def read(self, n=-1):
        return self._buf.read(n)


def test_read_workbook_skips_blank_rows_and_records_row_numbers():
    buf = _xlsx([_row(name='A'), {}, _row(name='B')])
    rows = ci.read_workbook(_Upload(buf))
    assert [r['name'] for r in rows] == ['A', 'B']
    assert [r['_row'] for r in rows] == [2, 4]


def test_read_workbook_rejects_missing_required_header():
    buf = _xlsx([{'name': 'A'}], headers=['name', 'phone'])
    with pytest.raises(ci.ImportFileError, match='address'):
        ci.read_workbook(_Upload(buf))


def test_read_workbook_rejects_non_xlsx():
    with pytest.raises(ci.ImportFileError):
        ci.read_workbook(_Upload(io.BytesIO(b'name,phone\nA,1\n'), 'c.csv'))


def test_read_workbook_rejects_empty_sheet():
    with pytest.raises(ci.ImportFileError, match='no data rows'):
        ci.read_workbook(_Upload(_xlsx([])))
